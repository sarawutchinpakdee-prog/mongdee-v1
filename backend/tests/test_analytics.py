import io
import tempfile
import time
import unittest
from datetime import date, datetime, timedelta
from pathlib import Path
from unittest.mock import patch

from app import config, db


DAY = 86400


class AnalyticsTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        base = Path(self.tmp.name)
        self.patch = patch.multiple(config, DATA_DIR=base, CAPTURES_DIR=base / "captures", DB_PATH=base / "test.db")
        self.patch.start()
        db.init_db()
        self.a = db.create_product("A", None, None, None, None, None, None, None)
        self.b = db.create_product("B", None, None, None, None, None, None, None)
        self.now = time.time()

    def tearDown(self):
        self.patch.stop()
        self.tmp.cleanup()

    def event(self, product_id, matched, age_days=0.0, corrected=False, confidence=0.7):
        with db.get_conn() as conn:
            conn.execute(
                "INSERT INTO scan_events (product_id, confidence, matched, corrected, created_at) VALUES (?, ?, ?, ?, ?)",
                (product_id, confidence, int(matched), int(corrected), self.now - age_days * DAY),
            )

    def test_every_chart_counts_only_matches_to_existing_products(self):
        self.event(self.a, True)
        self.event(self.a, True)
        self.event(self.b, True)
        self.event(None, True)   # product deleted later -> must not be counted anywhere
        self.event(None, False)  # unknown scan
        summary = {r["name"]: r["scan_count"] for r in db.analytics_summary()}
        self.assertEqual(summary, {"A": 2, "B": 1})
        self.assertEqual(sum(r["count"] for r in db.analytics_hourly()), 3)
        self.assertEqual(sum(r["count"] for r in db.analytics_daily()), 3)
        overview = db.analytics_overview()
        self.assertEqual((overview["successful"], overview["unmatched"], overview["orphaned"]), (3, 1, 1))

    def test_success_rate_and_corrections(self):
        for _ in range(3):
            self.event(self.a, True)
        self.event(self.a, True, corrected=True)
        self.event(None, False)
        overview = db.analytics_overview()
        self.assertEqual(overview["successful"], 4)
        self.assertEqual(overview["corrected"], 1)
        self.assertAlmostEqual(overview["success_rate"], 4 / 5)
        self.assertEqual((overview["products_total"], overview["products_scanned"]), (2, 1))

    def test_success_rate_is_none_without_any_attempts(self):
        self.assertIsNone(db.analytics_overview()["success_rate"])

    def test_ranges_filter_by_calendar_days(self):
        self.event(self.a, True, age_days=0)
        self.event(self.a, True, age_days=2)
        self.event(self.a, True, age_days=40)
        counts = {key: db.analytics_overview(key)["successful"] for key in ("today", "7d", "30d", "all")}
        self.assertEqual(counts, {"today": 1, "7d": 2, "30d": 2, "all": 3})
        self.assertEqual(db.analytics_summary("today")[0]["scan_count"], 1)
        self.assertEqual(sum(r["count"] for r in db.analytics_hourly("7d")), 2)

    def test_range_start_is_local_midnight(self):
        noon = datetime(2026, 9, 19, 12, 30)
        self.assertEqual(db.range_start("today", now=noon), datetime(2026, 9, 19).timestamp())
        self.assertEqual(db.range_start("7d", now=noon), datetime(2026, 9, 13).timestamp())
        self.assertEqual(db.range_start("all", now=noon), 0.0)

    def test_daily_fills_missing_days_with_zero(self):
        self.event(self.a, True, age_days=0)
        self.event(self.a, True, age_days=3)
        days = db.analytics_daily("7d")
        self.assertEqual(len(days), 7)
        self.assertEqual(days[-1]["day"], date.today().isoformat())
        self.assertEqual(sum(d["count"] for d in days), 2)
        self.assertEqual([d["count"] for d in days if d["count"] == 0].__len__(), 5)

    def test_daily_for_all_time_is_capped_to_sixty_days(self):
        self.event(self.a, True, age_days=200)
        self.event(self.a, True, age_days=1)
        days = db.analytics_daily("all")
        self.assertEqual(len(days), 60)
        self.assertEqual(sum(d["count"] for d in days), 1, "events older than the window are not drawn")

    def test_products_with_no_scans_are_still_listed(self):
        self.event(self.a, True)
        rows = db.analytics_summary("today")
        self.assertEqual([(r["name"], r["scan_count"]) for r in rows], [("A", 1), ("B", 0)])

    def test_new_day_reset_restarts_today_without_deleting_anything(self):
        self.event(self.a, True, age_days=30 / DAY)
        self.event(self.a, True, age_days=10 / DAY)
        db.set_day_reset(self.now - 20)
        self.assertEqual(db.analytics_overview("today")["successful"], 1)
        self.assertEqual(sum(r["scan_count"] for r in db.analytics_summary("today")), 1)
        self.assertEqual(sum(r["count"] for r in db.analytics_hourly("today")), 1)
        self.assertEqual(db.analytics_daily("today")[0]["count"], 1)
        self.assertEqual(db.analytics_overview("7d")["successful"], 2, "other periods still include everything")
        self.assertEqual(db.analytics_overview("all")["successful"], 2)
        self.assertTrue(db.analytics_overview("today")["day_reset"])

        db.clear_day_reset()
        self.assertEqual(db.analytics_overview("today")["successful"], 2)
        self.assertFalse(db.analytics_overview("today")["day_reset"])

    def test_a_reset_from_an_earlier_day_or_the_future_is_ignored(self):
        self.event(self.a, True, age_days=10 / DAY)
        for stamp in (self.now - 2 * DAY, self.now + 3600):
            db.set_day_reset(stamp)
            self.assertEqual(db.analytics_overview("today")["successful"], 1)
            self.assertFalse(db.day_reset_active())

    def test_export_workbook_has_every_sheet_and_the_same_numbers(self):
        from openpyxl import load_workbook
        from app import report_export
        self.event(self.a, True, confidence=0.8)
        self.event(self.a, True, corrected=True)
        self.event(None, False)
        self.event(None, True)  # product deleted later
        workbook = load_workbook(io.BytesIO(report_export.build_report("all")))
        self.assertEqual(workbook.sheetnames, ["สรุป", "สินค้า", "รายวัน", "รายชั่วโมง", "รายการสแกน"])

        summary = {row[0]: row[1] for row in workbook["สรุป"].iter_rows(values_only=True) if row[0]}
        self.assertEqual(summary["สแกนสำเร็จ (ครั้ง)"], 2)
        self.assertEqual(summary["สแกนไม่รู้จัก / ไม่แน่ใจ (ครั้ง)"], 1)
        self.assertEqual(summary["พนักงานแก้ผลสแกนเอง (ครั้ง)"], 1)
        self.assertEqual(summary["ไม่นับ: สแกนของสินค้าที่ถูกลบแล้ว (ครั้ง)"], 1)
        self.assertAlmostEqual(summary["อัตราสแกนสำเร็จ"], 2 / 3)

        products = list(workbook["สินค้า"].iter_rows(min_row=2, values_only=True))
        self.assertEqual([(p[0], p[3]) for p in products], [("A", 2), ("B", 0)])
        self.assertEqual(len(list(workbook["รายชั่วโมง"].iter_rows(min_row=2))), 24)
        results = [row[2] for row in workbook["รายการสแกน"].iter_rows(min_row=2, values_only=True)]
        self.assertEqual(sorted(results), sorted(["สำเร็จ", "สำเร็จ", "ไม่รู้จัก / ไม่แน่ใจ", "สำเร็จ (สินค้าถูกลบแล้ว)"]))


if __name__ == "__main__":
    unittest.main()
